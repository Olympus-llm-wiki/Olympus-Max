"""Validate the supported blank schedule and map identifiers to rows."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
import zipfile

from .extract import WorkflowError
from .ned import ASSUMPTIONS, SYSTEM, choose, normalize, system_fields


def inspect_template(path: Path) -> dict:
    import openpyxl
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 32 * 1024 * 1024:
            raise WorkflowError("unsupported_template", {"reason": "expanded_size_limit"})
        forbidden = [name for name in archive.namelist() if any(part in name.lower() for part in ("vbaproject", "externallink", "xl/drawings/", "xl/charts/", "xl/tables/", "comments"))]
        if forbidden:
            raise WorkflowError("unsupported_template", {"reason": "unvalidated_native_features", "parts": forbidden[:8]})
    book = openpyxl.load_workbook(path, data_only=False)
    if len(book.worksheets) != 1:
        raise WorkflowError("unsupported_template", {"reason": "one_sheet_required"})
    sheet = book.active
    headers = {"A2": "№ системы", "G2": "Вентилятор", "N2": "Воздухонагреватель водяной", "R2": "Воздухонагреватель электрический", "V2": "Воздухоохладитель", "AG2": "Фильтр",
               "E3": "Тип", "F3": "Производитель", "G3": "Расход", "J3": "Напор", "K3": "Мощность", "L3": "Парам, сети",
               "P3": "Расход тепла", "Q3": "Теплоноситель", "T3": "Мощность", "U3": "Парам, сети",
               "X3": "Расход холода", "Y3": "Холодоноситель", "Z3": "Компрессорно-конденсаторный блок"}
    header_key = lambda value: re.sub(r"[^a-zа-яё0-9]", "", normalize(str(value)).lower())
    mismatches = [cell for cell, text in headers.items() if header_key(sheet[cell].value) != header_key(text)]
    for cell in ("K5", "P5", "T5", "X5"):
        if normalize(str(sheet[cell].value)) != "кВт":
            mismatches.append(cell)
    if normalize(str(sheet["J5"].value)) != "Па":
        mismatches.append("J5")
    for cell in ("N5", "O5", "R5", "S5", "V5", "W5"):
        if normalize(str(sheet[cell].value)) not in ("C", "С", "°C"):
            mismatches.append(cell)
    if normalize(str(sheet["G5"].value)) not in ("м2/ч", "м3/ч"):
        mismatches.append("G5")
    if mismatches:
        raise WorkflowError("unsupported_template", {"reason": "header_mismatch", "cells": mismatches})
    if sheet.max_row > 1000 or sheet.max_column > 36:
        raise WorkflowError("unsupported_template", {"reason": "dimensions"})
    if any(cell.data_type == "f" for row in sheet for cell in row):
        raise WorkflowError("unsupported_template", {"reason": "formula_template_not_supported"})
    systems = []
    seen = set()
    last = 7
    for row in range(8, sheet.max_row + 1):
        value = sheet.cell(row, 1).value
        if value is None:
            continue
        name = normalize(str(value)).upper()
        if not SYSTEM.fullmatch(name):
            raise WorkflowError("unsupported_template", {"reason": "non_system_row", "row": row})
        if name in seen:
            raise WorkflowError("unsupported_template", {"reason": "duplicate_system", "system": name})
        seen.add(name)
        span = next((m.max_row - m.min_row + 1 for m in sheet.merged_cells.ranges if m.min_col == m.max_col == 1 and m.min_row == row), 1)
        if span not in (1, 2):
            raise WorkflowError("unsupported_template", {"reason": "system_row_span", "row": row})
        rows = list(range(row, row + span))
        systems.append({"system": name, "rows": rows})
        last = max(last, rows[-1])
    if not systems:
        raise WorkflowError("unsupported_template", {"reason": "no_system_rows"})
    if any(c.value is not None for row in sheet.iter_rows(min_row=last+1) for c in row):
        raise WorkflowError("unsupported_template", {"reason": "content_below_systems"})
    allowed = []
    for system in systems:
        for row in system["rows"]:
            for col in range(5, 37):
                cell = sheet.cell(row, col)
                if isinstance(cell, openpyxl.cell.cell.MergedCell):
                    continue
                if cell.value is not None:
                    raise WorkflowError("unsupported_template", {"reason": "parameter_area_not_blank", "cell": cell.coordinate})
                allowed.append(cell.coordinate)
    full_text = "\n".join(f"{sheet.title}!{c.coordinate}: {c.value}" for row in sheet for c in row if c.value is not None)
    # Notes go after the complete template, including unused formatted rows.
    # This prevents new full-width notes from intersecting pre-existing merges.
    return {"sheet": sheet.title, "systems": systems, "last_row": max(last, sheet.max_row), "allowed_cells": allowed, "full_text": full_text}


def map_schedule(parsed: dict, template: dict) -> dict:
    issues = deepcopy(parsed["issues"])
    updates = []
    power = []
    systems = []
    source_by_name = {s["system"]: s for s in parsed["systems"]}
    for record in template["systems"]:
        name, rows = record["system"], record["rows"]
        source = source_by_name.get(name)
        if source is None:
            issues.append({"code": "template_system_missing_from_pdf", "system": name})
            systems.append({**record, "pages": [], "directions": []})
            continue
        directions = {b["direction"] for b in source["blocks"].values() if b["kind"] == "fan"}
        if len(directions) == 2 and len(rows) != 2:
            raise WorkflowError("unsupported_template", {"reason": "two_directions_need_two_rows", "system": name})
        if len(rows) == 2 and directions != {"supply", "exhaust"}:
            raise WorkflowError("unsupported_template", {"reason": "two_rows_need_both_directions", "system": name})
        order = ["supply", "exhaust"] if len(rows) == 2 else [next(iter(directions), "supply")]
        systems.append({**record, "pages": source["pages"], "directions": [{"row": row, "direction": d} for row, d in zip(rows, order)]})

        def add(cell, value):
            if cell not in template["allowed_cells"]:
                issues.append({"code": "merged_or_unmapped_target", "system": name, "cell": cell})
                return
            ev = value["evidence"][0]
            updates.append({"cell": cell, **value, "page": ev["page"], "quote": ev["quote"], "bbox": ev["bbox"]})

        model = choose(source["model_candidates"], {"system": name, "field": "E"}, issues)
        if model:
            add(f"E{rows[0]}", model)
            add(f"F{rows[0]}", {**deepcopy(model), "value": "NED", "basis": "recognized_ned_profile"})
        else:
            issues.append({"code": "missing_model", "system": name})
        for row, direction in zip(rows, order):
            fields = system_fields(source, direction, parsed["source_sha256"], issues)
            for col, value in fields.items():
                if col != "T_installed":
                    add(f"{col}{row}", value)
            if "T" in fields or "T_installed" in fields:
                if "T_installed" not in fields:
                    issues.append({"code": "missing_installed_heater_power", "system": name, "row": row})
                power.append({"system": name, "cell": f"T{row}",
                              "consumed_kw": fields.get("T", {}).get("value"),
                              "installed_kw": fields.get("T_installed", {}).get("value"),
                              "page": (fields.get("T") or fields["T_installed"])["evidence"][0]["page"]})
    for source in parsed["systems"]:
        if source["system"] not in {s["system"] for s in systems}:
            issues.append({"code": "pdf_system_missing_from_template", "system": source["system"], "pages": source["pages"]})
    filled = {u["cell"] for u in updates}
    unfilled = [{"cells": [cell], "reason": "not_extracted", "explanation": "Не извлечено однозначно поддержанным профилем; не равно нулю и не доказывает отсутствие оборудования."}
                for cell in template["allowed_cells"] if cell not in filled]
    return {"schema": 1, "profile": parsed["profile"], "source_sha256": parsed["source_sha256"],
            "sheet": template["sheet"], "last_row": template["last_row"], "systems": systems,
            "updates": updates, "electric_power": power, "unfilled": unfilled,
            "assumptions": deepcopy(ASSUMPTIONS), "issues": issues,
            "status": "needs_review" if issues else "draft_with_assumptions"}
