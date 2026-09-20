"""Synthetic behavioral fixtures: no client documents or values are checked in."""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from olympus.khvs.extract import WorkflowError, extract_pdf, lines_from_words, write_json
from olympus.khvs.ned import parse_ned, system_fields
from olympus.khvs.template import inspect_template, map_schedule
from olympus.khvs.pipeline import in_git, render_page, resolve_input, run
from olympus.preservation import Store

HAVE_DOCUMENTS = all(importlib.util.find_spec(name) for name in ("pdfplumber", "pypdf", "reportlab", "openpyxl", "lxml"))


def line(text, y, side="left"):
    x = 55 if side == "left" else 310
    return {"text": text, "bbox": [x, y, x + 230, y + 9]}


def synthetic_page(number=1, *, name="П81", flow=1728, pressure=250, power=2.7):
    text = f"Номер коммерческого предложения ND99-987654/1\nНаименование установки {name} + бак.сек.\nAIRNED-M2\nТип установки TEST-1 [Напольная]\nДата коммерческого предложения 01.01.2020"
    left = [line("ВЕНТИЛЯТОР", 170), line(f"Расход воздуха (м3/ч) {flow}", 190),
            line(f"P свободное (Па) {pressure}", 210), line("Количество агрегатов (шт) 1", 230),
            line("НАГРЕВАТЕЛЬ 1", 280), line("Мощность нагрева потребляемая (кВт) 6.4", 300),
            line("Мощность нагрева установочная (кВт) 9", 320), line("Напряжение/Число ступеней 400 / 2", 340),
            line("t°/влажность вх. воздуха (°C / %) 11", 360), line("t°/влажность вых. воздуха (°C / %) 24", 380),
            line("ФИЛЬТР СТУПЕНЬ 1", 440), line("Класс очистки EU3", 460), line("АКУСТИЧЕСКИЕ ХАРАКТЕРИСТИКИ", 600)]
    right = [line(f"Номинальная мощность (Nуст, кВт) {power}", 190, "right"), line("Напряжение (В) 400", 210, "right"),
             line("НАГРЕВАТЕЛЬ 2", 280, "right"), line("Мощность нагрева потребляемая (кВт) 35.7", 300, "right"),
             line("t°/влажность вх. воздуха (°C / %) -15", 320, "right"), line("t°/влажность вых. воздуха (°C / %) 24", 340, "right"),
             line("Тип теплоносителя WTR", 360, "right"), line("t° вх. жидкости (°C) 80", 380, "right"),
             line("t° вых. жидкости (°C) 60", 400, "right"), line("ФИЛЬТР СТУПЕНЬ 2", 440, "right"), line("Класс очистки EU7", 460, "right")]
    return {"number": number, "width": 595, "height": 842, "text": text, "lines": [line("Приточная часть", 150)], "left": left, "right": right}


def extraction(*pages):
    return {"source_sha256": "a" * 64, "pages": list(pages)}


class NedProfileTests(unittest.TestCase):
    def test_side_by_side_heaters_and_wrapped_model(self):
        data = parse_ned(extraction(synthetic_page()))
        self.assertEqual(data["issues"], [])
        self.assertEqual(data["systems"][0]["model_candidates"][0]["value"], "AIRNED-M2 TEST-1 [Напольная]")
        issues = []
        fields = system_fields(data["systems"][0], "supply", "a" * 64, issues)
        self.assertEqual(issues, [])
        self.assertEqual([fields[c]["value"] for c in ("G", "J", "K", "L")], [1728, 250, 2.7, 400])
        self.assertEqual([fields[c]["value"] for c in ("P", "T", "T_installed")], [35.7, 6.4, 9])
        self.assertEqual(fields["Q"]["value"], "WTR; 80/60 °C")
        self.assertEqual(fields["T"]["assumption_id"], "Д1")
        self.assertIn("Бактерицидная", fields["AJ"]["value"])

    def test_conflict_never_chooses_first(self):
        page = synthetic_page()
        page["left"].insert(2, line("Расход воздуха (м3/ч) 999", 200))
        data = parse_ned(extraction(page)); issues = []
        fields = system_fields(data["systems"][0], "supply", "a" * 64, issues)
        self.assertNotIn("G", fields)
        self.assertTrue(any(x["code"] == "conflicting_values" for x in issues))

    def test_unexpected_unit_and_malformed_number(self):
        for text in ("Расход воздуха (л/с) 1728", "Расход воздуха (м3/ч) 55..111"):
            page = synthetic_page(); page["left"][1]["text"] = text
            data = parse_ned(extraction(page)); issues = []
            fields = system_fields(data["systems"][0], "supply", "a" * 64, issues)
            self.assertNotIn("G", fields)
            self.assertTrue(issues)

    def test_exhaust_continues_on_next_page(self):
        page = synthetic_page(name="ПВ81")
        page["lines"].append(line("Вытяжная часть", 690))
        page["left"].extend([line("ВЕНТИЛЯТОР", 710), line("Расход воздуха (м3/ч) 555", 735)])
        page["right"].append(line("Установочная мощность (Nуст, кВт) 0.4", 735, "right"))
        tail = {"number": 2, "width": 595, "height": 842, "text": "Продолжение", "lines": [line("Вытяжная часть", 5)],
                "left": [line("P свободное (Па) 175", 20), line("АКУСТИЧЕСКИЕ ХАРАКТЕРИСТИКИ", 100)],
                "right": [line("Напряжение (В) 230", 20, "right")]}
        data = parse_ned(extraction(page, tail)); issues = []
        fields = system_fields(data["systems"][0], "exhaust", "a" * 64, issues)
        self.assertEqual([fields[c]["value"] for c in ("G", "J", "K", "L")], [555, 175, .4, 230])
        self.assertEqual(fields["J"]["evidence"][0]["page"], 2)
        self.assertNotIn("AG", fields)

    def test_missing_present_component_parameter_requires_review(self):
        page = synthetic_page(); page["left"] = [x for x in page["left"] if not x["text"].startswith("Напряжение/")]
        data = parse_ned(extraction(page)); issues = []
        fields = system_fields(data["systems"][0], "supply", "a" * 64, issues)
        self.assertNotIn("U", fields)
        self.assertTrue(any(x["code"] == "missing_component_field" and x["field"] == "U" for x in issues))

    def test_foreign_profile_and_empty_text(self):
        for text in ("", "Русклимат. Название: П1"):
            page = synthetic_page(); page["text"] = text
            with self.assertRaisesRegex(WorkflowError, "needs_profile"):
                parse_ned(extraction(page))


@unittest.skipUnless(HAVE_DOCUMENTS, "Optional KHVS document dependencies are not installed")
class DocumentWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def make_pdf(self, path):
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(path)); c.drawString(40, 700, "Synthetic text source"); c.save()

    def make_template(self, names):
        import openpyxl
        from openpyxl.styles import Border, Side
        b = openpyxl.Workbook(); s = b.active; s.title = "Table 1"
        headers = {"A2": "№ системы", "G2": "Вентилятор", "N2": "Воздухонагреватель водяной", "R2": "Воздухонагреватель электрический", "V2": "Воздухоохладитель", "AG2": "Фильтр",
                   "E3": "Тип", "F3": "Произво-дитель", "G3": "Расход", "J3": "Напор", "K3": "Мощность", "L3": "Парам, сети",
                   "P3": "Расход тепла", "Q3": "Теплоноситель", "T3": "Мощность", "U3": "Парам, сети", "X3": "Расход холода", "Y3": "Холодо-носитель", "Z3": "Компрессорно-конденсаторный блок",
                   "G5": "м3/ч", "J5": "Па", "K5": "кВт", "P5": "кВт", "T5": "кВт", "X5": "кВт",
                   "N5": "С", "O5": "С", "R5": "С", "S5": "С", "V5": "С", "W5": "С"}
        for addr, value in headers.items(): s[addr] = value
        for row, name in enumerate(names, 8):
            s.cell(row, 1, name); s.cell(row, 2, "Синтетическое помещение"); s.cell(row, 4, 1)
            for col in range(1, 37):s.cell(row, col).border = Border(bottom=Side(style="thin"))
            s.merge_cells(start_row=row, start_column=7, end_row=row, end_column=9)
        s.merge_cells("A1:AJ1"); s["A1"] = "Синтетический шаблон"
        path = self.root / "template.xlsx"; b.save(path); return path

    def test_cache_hit_and_corruption_rebuild(self):
        path = self.root / "source.pdf"; self.make_pdf(path)
        first, cold = extract_pdf(path, self.root / "cache")
        with patch("pdfplumber.open", side_effect=AssertionError("cache should skip extraction")):
            second, warm = extract_pdf(path, self.root / "cache")
        self.assertEqual(first, second); self.assertEqual(cold["extraction_calls"], 1); self.assertEqual(warm["extraction_calls"], 0)
        cache = self.root / "cache" / (cold["key"] + ".json")
        envelope = json.loads(cache.read_bytes()); envelope["payload"]["pages"][0]["text"] = "corrupted"; cache.write_text(json.dumps(envelope))
        rebuilt, info = extract_pdf(path, self.root / "cache")
        self.assertEqual(info["state"], "rebuilt_corrupt"); self.assertEqual(rebuilt, first)

    def test_reordered_template_and_input_preservation(self):
        path = self.make_template(["П82", "П81"]); before = path.read_bytes()
        parsed = parse_ned(extraction(synthetic_page(name="П81", flow=111), synthetic_page(2, name="П82", flow=222)))
        data = map_schedule(parsed, inspect_template(path)); values = {u["cell"]: u["value"] for u in data["updates"]}
        self.assertEqual(values["G8"], 222); self.assertEqual(values["G9"], 111)
        self.assertEqual(path.read_bytes(), before)

    def test_existing_parameter_is_not_overwritten(self):
        import openpyxl
        path = self.make_template(["П81"]); b = openpyxl.load_workbook(path); b.active["K8"] = 123; b.save(path)
        with self.assertRaisesRegex(WorkflowError, "unsupported_template"):
            inspect_template(path)

    def test_duplicate_system_and_formula_rejected(self):
        import openpyxl
        path = self.make_template(["П81", "П81"])
        with self.assertRaisesRegex(WorkflowError, "unsupported_template"):inspect_template(path)
        path = self.make_template(["П81"]); b = openpyxl.load_workbook(path); b.active["B8"] = '=1+1'; b.save(path)
        with self.assertRaisesRegex(WorkflowError, "unsupported_template"):inspect_template(path)

    def test_wrong_target_unit_is_not_silently_relabelled(self):
        import openpyxl
        path = self.make_template(["П81"]); b = openpyxl.load_workbook(path); b.active["T5"] = "Вт"; b.save(path)
        with self.assertRaisesRegex(WorkflowError, "unsupported_template"):inspect_template(path)

    def test_failed_profile_keeps_archive_and_page_view(self):
        path = self.root / "source.pdf"; self.make_pdf(path); template = self.make_template(["П81"])
        out, state = self.root / "job", self.root / "state"
        with self.assertRaisesRegex(WorkflowError, "needs_profile"):
            run(path, template, out, state)
        manifest = json.loads((out / "manifest.json").read_bytes())
        self.assertEqual(manifest["status"], "failed")
        self.assertTrue(all(x["memory"] == "archived" for x in manifest["receipts"].values()))
        rendered = render_page(out, 1)
        self.assertTrue(Path(rendered["image"]).is_file())
        (out / "inputs/source.pdf").write_bytes(b"altered")
        with self.assertRaisesRegex(WorkflowError, "source_snapshot_mismatch"):render_page(out, 1)

    def test_destination_and_git_guards(self):
        path = self.root / "source.pdf"; self.make_pdf(path); template = self.make_template(["П81"])
        out = self.root / "existing"; out.mkdir(); marker = out / "user.txt"; marker.write_text("keep")
        with self.assertRaisesRegex(WorkflowError, "output_directory_not_empty"):run(path, template, out, self.root / "state")
        self.assertEqual(marker.read_text(), "keep")
        repo = self.root / "repo"; repo.mkdir(); (repo / ".git").mkdir()
        with self.assertRaisesRegex(WorkflowError, "private_documents_require_non_git_paths"):run(path, template, repo / "out", self.root / "state")

    def test_complete_synthetic_pipeline_archives_without_models(self):
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from olympus.khvs.workbook import runtime_paths
        paths = runtime_paths()
        font = paths["modules"] / "pdfjs-dist/standard_fonts/LiberationSans-Regular.ttf"
        if not paths["node"].is_file() or not font.is_file() or not (paths["modules"] / "@oai/artifact-tool").is_dir():
            self.skipTest("Artifact runtime unavailable")
        pdfmetrics.registerFont(TTFont("SyntheticFixture", str(font)))
        page = synthetic_page(); source = self.root / "source.pdf"
        c = canvas.Canvas(str(source), pagesize=(595, 842)); c.setFont("SyntheticFixture", 8)
        for i, text in enumerate(page["text"].splitlines()):c.drawString(45, 810-i*12, text)
        for row in page["lines"] + page["left"] + page["right"]:
            c.drawString(row["bbox"][0], 842-row["bbox"][1]-8, row["text"])
        c.showPage(); c.showPage(); c.save()  # A blank source page must survive the review PDF.
        template = self.make_template(["П81"])
        result = run(source, template, self.root / "job", self.root / "state", preview=False)
        self.assertEqual(result["status"], "draft_with_assumptions")
        self.assertEqual(result["model_calls"], 0)
        data = json.loads((self.root / "job/values.json").read_bytes())
        values = {u["cell"]: u["value"] for u in data["updates"]}
        self.assertEqual(values["T8"], 6.4); self.assertEqual(values["P8"], 35.7)
        self.assertTrue((self.root / "job/package.zip").is_file())
        states = Store(self.root / "state").status()["delivery"]
        self.assertGreater(states.get("archived", 0), 0); self.assertEqual(states.get("pending", 0), 0)


if __name__ == "__main__":
    unittest.main()
